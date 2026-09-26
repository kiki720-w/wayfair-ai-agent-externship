from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / ".docx-build" / "Wayfair_Extern_Prompt_Generator_Submission.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=11, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True)
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
styles["Body Text"].font.name = "Aptos"
styles["Body Text"].font.size = Pt(11)
styles["Body Text"].paragraph_format.space_after = Pt(7)
styles["Body Text"].paragraph_format.line_spacing = 1.08
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(26)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run_font(title.add_run("Wayfair Extern Prompt Generator Submission"), name="Aptos Display", size=26, bold=True)

subtitle = doc.add_paragraph()
set_run_font(subtitle.add_run("n8n AI Agent Engineering Externship"), size=12, bold=True, color="3C4858")
meta = doc.add_paragraph()
set_run_font(meta.add_run("Prepared 27 September 2026"), size=10, color="5F6B7A")

add_body(
    doc,
    "This submission documents a working n8n prompt-generation agent that converts a short rug design idea into a structured moodboard prompt. The complete workflow was tested end to end with Gemini 3.8 Flash and returned a detailed, image-generation-ready output.",
)

add_heading(doc, "Submission Summary", 1)
table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Inches(2.0)
table.columns[1].width = Inches(4.8)
headers = table.rows[0].cells
headers[0].text = "Requirement"
headers[1].text = "Submission Detail"
set_repeat_table_header(table.rows[0])
for cell in headers:
    set_cell_shading(cell, "26364A")
    set_cell_borders(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, bold=True, color="FFFFFF")

rows = [
    ("Workflow built", "Yes"),
    ("Workflow", "On Chat Message -> AI Agent -> Google Gemini Chat Model"),
    ("Model", "Gemini 3.8 Flash"),
    ("Verified run", "Success in 6.443 seconds with 459 output tokens"),
]
for idx, (left, right) in enumerate(rows):
    cells = table.add_row().cells
    cells[0].text = left
    cells[1].text = right
    for cell in cells:
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if idx % 2 == 1:
            set_cell_shading(cell, "F3F6F9")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, bold=(cell is cells[0]))

doc.add_paragraph()
add_heading(doc, "Style Idea", 1)
add_body(
    doc,
    "Design a cozy Scandinavian living room rug moodboard using warm ivory, oatmeal, and muted sage, with natural wool and subtle geometric texture.",
)

add_heading(doc, "Generated Prompt", 1)
add_body(
    doc,
    "A professional interior design moodboard centered on a cozy Scandinavian living room area rug as the hero subject, presented in a clean and cohesive catalog-style layout. The rug features an understated, modern geometric high-low pattern woven from plush natural wool in a soothing palette of warm ivory, heathered oatmeal, and soft muted sage green. The board includes a styled Nordic living room scene showcasing the rug anchored on light white-oak flooring beneath a low minimalist oak coffee table and raw linen sofa corner, illuminated by soft, diffused morning daylight. Flanking panels display tactile macro close-ups highlighting the hand-tufted loop pile texture, yarn twist, and clean serged edge details, alongside complementary decor elements including ribbed matte sage ceramic vessels, an unbleached linen fabric swatch, and light blonde wood samples. Shot with crisp, medium-format editorial photography, balanced lighting, gentle realistic shadows, and an inviting, serene hygge atmosphere, 8k resolution. Negative constraints: no people, no text, no logos, no watermarks, no distorted furniture, no duplicate objects.",
)

add_heading(doc, "Workflow Evidence", 1)
add_bullet(doc, "The chat trigger received the style idea successfully.")
add_bullet(doc, "The AI Agent called the connected Gemini chat model.")
add_bullet(doc, "All three nodes completed successfully in the verified execution.")
add_bullet(doc, "The final output preserved the requested subject, palette, material, style, room context, photography direction, and negative constraints.")

add_heading(doc, "Reflection", 1)
add_body(
    doc,
    "I wanted to turn a short and ambiguous rug-style idea into a prompt detailed enough to create a consistent, presentation-ready visual direction. The main improvement came from separating the subject, style, material, color palette, room context, composition, lighting, and negative constraints. Compared with a short prompt, the structured result gives the image model clearer control over what must remain prominent - the rug - and what should only support it. Constraints such as no text, no people, and no duplicate objects also reduce common visual-generation failures. The workflow demonstrates how a reusable system prompt can standardize creative input before it reaches an image-generation model.",
)

add_heading(doc, "Implementation Note", 1)
add_body(
    doc,
    "The course template originally referenced Gemini 2.5 Flash. Google returned a 404 stating that the model was no longer available to new users, so I migrated the workflow to Gemini 3.8 Flash. The first request encountered a temporary high-demand 503 response; retrying succeeded without changing the prompt or workflow logic. Credentials are stored only in the local n8n credential store and are excluded from the GitHub repository.",
)

add_heading(doc, "Portfolio Archive", 1)
add_body(doc, "GitHub repository: https://github.com/kiki720-w/wayfair-ai-agent-externship")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
